"""Данные для экрана: весь SQL приложения только здесь (docs/partner_tasks.md, раздел 2), плюс подготовка данных к показу.

load_results() открывает базу только на чтение. Расчётов расхождений здесь нет: числа берутся из views position_status, issues_view и
run_summary, которые строит run_pipeline; здесь только форматирование, сборка фраз, сортировка, фильтры и подсчёты для диаграмм.
Всё, что не SQL (fmt_num, build_cards, filter_issues, chart_frames, ...), работает с датафреймами и не знает про Streamlit.
"""
import json
import re
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

import pandas as pd

from src.config import load_config

TYPE_ORDER = ["volume_exceeded", "price_increase", "missing_in_vor", "late_act"]
TYPE_RU = {"volume_exceeded": "Превышение объёма", "price_increase": "Рост цены", "missing_in_vor": "Нет в ВОР",
           "late_act": "Акт после срока"}
TYPE_CHART_RU = {**TYPE_RU, "late_act": "Акт после срока договора"}
SEVERITY_ORDER = ["high", "medium", "low"]
SEVERITY_RU = {"high": "высокая", "medium": "средняя", "low": "низкая"}
SEVERITY_COLOR = {"high": "#C62828", "medium": "#F9A825", "low": "#9AA3AE"}
# цвет заливки и цвет текста: контраст не ниже 4.5:1 (проверяется тестом)
STATUS_COLORS = {"red": ("#C62828", "#FFFFFF"), "yellow": ("#F9A825", "#1B1B1B"), "green": ("#2E7D32", "#FFFFFF")}
STATUS_ROW_FILL = {"red": "#FDECEA", "yellow": "#FFF6DB", "green": "#E8F5E9"}       # светлые заливки строк таблицы, текст тёмный
STATUS_PLURAL = {"red": "Красные", "yellow": "Жёлтые", "green": "Зелёные"}
STATUS_SINGULAR = {"red": "Красный", "yellow": "Жёлтый", "green": "Зелёный"}
DOC_TYPE_RU = {"vor": "Ведомость объёмов работ (ВОР)", "estimate": "Смета", "contract": "Договор", "act": "Акт выполненных работ"}
TEMPLATE_RU = {"vor_a": "ВОР, каркас А", "vor_b": "ВОР, каркас Б", "estimate_a": "Смета", "contract_a": "Договор",
               "act_a": "Акт, форма А", "act_b": "Акт, форма Б", "act_c": "Акт, форма В"}
PREVIEW_ROWS = 80
SOURCE_ROLE = {"vor": "ВОР", "estimate": "Смета", "act": "Акт", "contract": "Договор"}
LOW_CONFIDENCE_NOTE = "Низкая уверенность, требует проверки"
AI_NOT_FOUND_NOTE = "ИИ не нашёл пару в ВОР, требует проверки"
EMPTY = "—"
ROOT = Path(__file__).resolve().parents[1]
MAX_AI_PAIRS_ON_CARD = 3


def connect_readonly(db_path) -> sqlite3.Connection:
    """Только чтение: запись в демо-базу через это соединение невозможна."""
    return sqlite3.connect(f"{Path(db_path).resolve().as_uri()}?mode=ro", uri=True)


# ---------- чтение из базы (единственное место с SQL) ----------
def load_results(db_path, project_id: str = "demo") -> dict:
    """-> {"summary": dict, "issues", "positions", "documents": DataFrame}.

    summary: json из run_summary последнего прогона проекта плюс banner; пустой словарь, если прогонов нет.
    issues: issues_view плюс колонки sources (строки «роль · файл · лист · строка» по документам, где нашлось расхождение) и
    ai_matched (пара названий склеена моделью), ai_pairs (что сопоставил ИИ: название в документе, в ВОР, уверенность, причина) и
    ai_none (ИИ не нашёл пару в ВОР: уверенность и причина). ai_examples: примеры разделения работы кода и ИИ. documents: файлы проекта и число непройденных проверок качества (диагностические
    записи «kind не определён» и «не сопоставлено, требует проверки» не считаются)."""
    cfg = load_config()
    diagnostic = (cfg["rules"]["matching"]["kind_unknown_dq_check"], cfg["rules"]["matching"]["ambiguous_dq_check"])
    conn = connect_readonly(db_path)
    try:
        row = conn.execute("SELECT summary_json, banner FROM run_summary WHERE project_id = ? ORDER BY run_id DESC LIMIT 1",
                           (project_id,)).fetchone()
        summary = {}
        if row:
            summary = json.loads(row[0])
            summary["banner"] = row[1]
        issues = pd.read_sql_query("SELECT * FROM issues_view WHERE project_id = ?", conn, params=(project_id,))
        positions = pd.read_sql_query("SELECT * FROM position_status WHERE project_id = ?", conn, params=(project_id,))
        items = pd.read_sql_query("SELECT doc_type, work_key, work_name_raw, quantity, unit_price, source_file, source_sheet, source_row "
                                  "FROM items WHERE project_id = ?", conn, params=(project_id,))
        ai_keys = {r[0] for r in conn.execute(
            "SELECT DISTINCT m.work_key FROM matches m JOIN staging_matches_ext e ON e.match_id = m.match_id "
            "JOIN items i ON i.item_id = m.item_id WHERE i.project_id = ? AND e.stage IN ('llm', 'llm_row')", (project_id,))}
        min_conf = cfg["rules"]["llm"]["min_confidence"]
        ai_pair_rows = pd.read_sql_query(
            "SELECT m.work_key, e.stage, e.reason, e.confidence, i.doc_type, i.work_name_raw AS doc_name, v.work_name_raw AS vor_name "
            "FROM matches m JOIN staging_matches_ext e ON e.match_id = m.match_id JOIN items i ON i.item_id = m.item_id "
            "JOIN items v ON v.item_id = m.matched_to_item_id WHERE i.project_id = ? AND e.stage IN ('llm', 'llm_row') ORDER BY m.match_id",
            conn, params=(project_id,))
        none_rows = pd.read_sql_query(
            "SELECT i.source_file, i.source_sheet, i.source_row, d.confidence, d.reason FROM staging_llm_row_decisions d "
            "JOIN staging_match_rows r ON r.group_id = d.group_id JOIN items i ON i.item_id = r.item_id "
            "WHERE d.project_id = ? AND d.outcome = 'none' AND d.confidence >= ?", conn, params=(project_id, min_conf))
        examples = collect_ai_examples(conn, project_id, cfg)
        files = pd.read_sql_query(
            "SELECT d.file_name, d.doc_type, d.status, e.template, e.currency, e.n_rows_read, e.n_unrecognized, e.n_service "
            "FROM documents d JOIN staging_documents_ext e ON e.doc_id = d.doc_id WHERE d.project_id = ? ORDER BY d.doc_id",
            conn, params=(project_id,))
        order = {"vor": 0, "estimate": 1, "contract": 2, "act": 3}                  # в списке сначала ВОР, потом смета, договор, акты
        files = files.assign(_o=files["doc_type"].map(order)).sort_values(["_o", "file_name"]).drop(columns="_o").reset_index(drop=True)
        marks = ",".join("?" for _ in diagnostic)
        documents = pd.read_sql_query(
            "SELECT d.doc_id, d.doc_type, d.file_name, d.status, "
            f"COUNT(CASE WHEN c.passed = 0 AND c.check_name NOT IN ({marks}) THEN 1 END) AS failed_checks "
            "FROM documents d LEFT JOIN dq_checks c ON c.doc_id = d.doc_id WHERE d.project_id = ? "
            "GROUP BY d.doc_id ORDER BY d.doc_id", conn, params=(*diagnostic, project_id))
    finally:
        conn.close()
    markers = cfg["synonyms"].get("strip_markers", [])
    issues = attach_sides(attach_ai(attach_context(issues, items, ai_keys), ai_pair_rows, none_rows, markers), items, markers)
    return {"summary": summary, "issues": issues, "positions": positions, "documents": documents, "ai_examples": pick_examples(examples),
            "ai_pairs_all": ai_pair_list(ai_pair_rows, markers), "files": files}


# ---------- форматирование ----------
def fmt_num(value) -> str:
    """Тысячи через пробел, десятичная запятая, не больше двух знаков; пусто -> «—»."""
    if value is None or pd.isna(value):
        return EMPTY
    text = f"{float(value):,.2f}".rstrip("0").rstrip(".")
    return text.replace(",", " ").replace(".", ",")


def fmt_pct(value, signed: bool = True) -> str:
    """Проценты с одним знаком после запятой: «+20,0%»."""
    if value is None or pd.isna(value):
        return EMPTY
    text = f"{abs(float(value)):.1f}".replace(".", ",")
    sign = ("+" if value >= 0 else "−") if signed else ""
    return f"{sign}{text}%"


def fmt_conf(value) -> str:
    """Уверенность ИИ: два знака после запятой («0,95»)."""
    return EMPTY if value is None or pd.isna(value) else f"{float(value):.2f}".replace(".", ",")


def fmt_date(iso) -> str:
    """«2025-09-30» -> «30.09.2025»; всё остальное возвращается как есть."""
    try:
        return date.fromisoformat(str(iso)).strftime("%d.%m.%Y")
    except ValueError:
        return str(iso)


def _luminance(hex_color: str) -> float:
    channels = [int(hex_color.lstrip("#")[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in channels]
    return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]


def contrast_ratio(color_a: str, color_b: str) -> float:
    """Контраст двух цветов по WCAG (нужно не меньше 4,5 для обычного текста)."""
    hi, lo = sorted((_luminance(color_a), _luminance(color_b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


# ---------- связанные строки и признак «ИИ помог» ----------
def attach_context(issues: pd.DataFrame, items: pd.DataFrame, ai_keys: set) -> pd.DataFrame:
    """Добавляет к расхождениям колонки sources (где смотреть: список строк «Роль · файл · лист «…» · строка N») и ai_matched.

    Источник расхождения всегда строка акта; к нему добавляются документы, с которыми сравнивали: ВОР и остальные строки актов по ключу
    для превышения объёма, смета для роста цены, договор для акта после срока."""
    issues = issues.copy()
    sources, ai = [], []
    for _, r in issues.iterrows():
        own = ("act", r["source_file"], r["source_sheet"], int(r["source_row"]))
        found = [own]
        same = items[items["work_key"] == r["work_key"]] if r["work_key"] else items.iloc[0:0]

        def pick(doc_type, only_first=True):
            part = same[same["doc_type"] == doc_type].sort_values(["source_file", "source_row"])
            part = part.head(1) if only_first else part
            return [(doc_type, p.source_file, p.source_sheet, int(p.source_row)) for p in part.itertuples()]

        if r["issue_type"] == "volume_exceeded":
            found += pick("vor") + [x for x in pick("act", False) if x != own]
        elif r["issue_type"] == "price_increase":
            found += pick("estimate")
        elif r["issue_type"] == "late_act":
            contract = items[items["doc_type"] == "contract"].head(1)
            found += [("contract", p.source_file, p.source_sheet, int(p.source_row)) for p in contract.itertuples()]
        sources.append([f"{SOURCE_ROLE[t]} · {f} · лист «{s}» · строка {n}" for t, f, s, n in found])
        ai.append(r["issue_type"] in ("volume_exceeded", "price_increase") and r["work_key"] in ai_keys)
    issues["sources"] = sources
    issues["ai_matched"] = ai
    return issues


# ---------- исходные файлы: каталог, структура, просмотр ----------
def file_label(row) -> str:
    """Подпись файла в списке выбора: «ВОР, каркас А · vor_1.xlsx»."""
    return f"{TEMPLATE_RU.get(row['template'], row['template'])} · {row['file_name']}"


def describe_file(row, templates: dict) -> dict:
    """Структура файла для панели: что это за документ, какой шаблон найден, какие колонки программа распознаёт, что прочитано.

    templates: config/templates.yaml (раздел templates). Колонки берутся из шаблона: первое название каждой колонки без служебных."""
    tpl = templates.get(row["template"], {})
    columns = [names[0] for field, names in tpl.get("columns", {}).items() if field != "seq" and names]
    if tpl.get("columns", {}).get("seq"):
        columns.insert(0, tpl["columns"]["seq"][0])
    label = "Колонки, которые распознаёт программа"
    if not columns and tpl.get("fields"):                       # договор: пары «подпись - значение», а не таблица
        columns = [names[0] for names in tpl["fields"].values() if names]
        label = "Подписи строк, которые распознаёт программа"
    return {"file": row["file_name"], "role": DOC_TYPE_RU.get(row["doc_type"], row["doc_type"]),
            "template": TEMPLATE_RU.get(row["template"], row["template"]), "columns": columns, "columns_label": label,
            "facts": [("Позиций прочитано", fmt_num(row["n_rows_read"])), ("Служебных строк", fmt_num(row["n_service"])),
                      ("Не распознано", fmt_num(row["n_unrecognized"])), ("Валюта", row["currency"] or EMPTY)]}


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return fmt_num(value)
    return str(value).strip()


def file_preview(path, max_rows: int = PREVIEW_ROWS) -> dict:
    """Первый лист Excel как в файле: номера строк, буквы колонок, значения и формулы (формулы показываются текстом «=…»).

    -> {sheet, n_rows (до последней непустой строки), n_cols, shown, frame}. Файл только читается."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path, data_only=False, read_only=True)
    ws = wb.worksheets[0]
    rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    title = ws.title
    wb.close()
    while rows and not any(rows[-1]):
        rows.pop()
    width = max((max((i for i, v in enumerate(r) if v), default=-1) + 1 for r in rows), default=0)
    rows = [(r + [""] * width)[:width] for r in rows]
    shown = rows[:max_rows]
    frame = pd.DataFrame(shown, columns=[get_column_letter(i + 1) for i in range(width)])
    frame.insert(0, "Строка", range(1, len(shown) + 1))
    return {"sheet": title, "n_rows": len(rows), "n_cols": width, "shown": len(shown), "frame": frame}


# ---------- «как написано»: что в ВОР и что в акте ----------
def attach_sides(issues: pd.DataFrame, items: pd.DataFrame, markers) -> pd.DataFrame:
    """Колонка sides: {left, right} для карточки «в ВОР и в акте, как написано».

    left: документ, с которым сравнивали (ВОР, смета или договор): название как в файле, значение. right: строки акта: название как в файле,
    значение, где лежит. Названия и числа берутся из items (название строки как записано, количество и цена после приведения единиц)."""
    sides = []
    for _, r in issues.iterrows():
        unit = r["unit"] or ""
        key = r["work_key"]
        same = items[items["work_key"] == key] if key else items.iloc[0:0]
        own = items[(items["source_file"] == r["source_file"]) & (items["source_sheet"] == r["source_sheet"])
                    & (items["source_row"] == r["source_row"]) & (items["doc_type"] == "act")]

        def row_of(part, value):
            return [{"name": _clean(x.work_name_raw, markers), "value": value(x), "where": f"{x.source_file} · строка {int(x.source_row)}"}
                    for x in part.itertuples()]

        kind = r["issue_type"]
        if kind == "volume_exceeded":
            vor = row_of(same[same["doc_type"] == "vor"].sort_values(["source_file", "source_row"]).head(1), lambda x: f"{fmt_num(r['expected'])} {unit}")
            acts = row_of(same[same["doc_type"] == "act"].sort_values(["source_file", "source_row"]), lambda x: f"{fmt_num(x.quantity)} {unit}")
            left = {"label": "В ВОР", "name": vor[0]["name"] if vor else None, "value": f"{fmt_num(r['expected'])} {unit}",
                    "where": vor[0]["where"] if vor else ""}
            right = {"label": "В актах", "rows": acts, "total": f"Всего {fmt_num(r['actual'])} {unit}"}
        elif kind == "price_increase":
            est = row_of(same[same["doc_type"] == "estimate"].head(1), lambda x: "")
            price = "цены в смете нет" if pd.isna(r["expected"]) else f"{fmt_num(r['expected'])} сом за {unit}"
            left = {"label": "В смете", "name": est[0]["name"] if est else None, "value": price, "where": est[0]["where"] if est else ""}
            right = {"label": "В акте", "rows": row_of(own, lambda x: f"{fmt_num(r['actual'])} сом за {unit}"), "total": ""}
        elif kind == "missing_in_vor":
            left = {"label": "В ВОР", "name": None, "value": "такой позиции не найдено", "where": ""}
            right = {"label": "В акте", "rows": row_of(own, lambda x: f"{fmt_num(x.quantity)} {unit}"), "total": ""}
        else:
            left = {"label": "Договор", "name": "Срок выполнения работ", "value": fmt_date(r["expected_text"]), "where": ""}
            right = {"label": "Акт", "rows": [{"name": f"Акт {r['source_file']}", "value": f"датирован {fmt_date(r['actual_text'])}",
                                               "where": f"{r['source_file']} · строка {int(r['source_row'])}"}],
                     "total": f"Позже срока на {fmt_num(r['actual'])} дн."}
        names = [x["name"].lower() for x in right["rows"]]
        left["names_differ"] = bool(left["name"] and kind in ("volume_exceeded", "price_increase") and any(n != left["name"].lower() for n in names))
        sides.append({"left": left, "right": right})
    issues = issues.copy()
    issues["sides"] = sides
    return issues


# ---------- что сделал ИИ ----------
def _clean(name, markers) -> str:
    text = str(name or "")
    for marker in markers:
        text = text.replace(marker, "")
    return re.sub(r"\s+", " ", text).strip()


def attach_ai(issues: pd.DataFrame, pair_rows: pd.DataFrame, none_rows: pd.DataFrame, markers) -> pd.DataFrame:
    """Колонки ai_pairs (пары названий «в документе -> в ВОР», которые сопоставил ИИ, по ключу позиции) и ai_none (ИИ не нашёл пару)."""
    by_key = {}
    for r in pair_rows.itertuples():
        pairs = by_key.setdefault(r.work_key, [])
        pair = {"doc_type": r.doc_type, "doc": _clean(r.doc_name, markers), "vor": _clean(r.vor_name, markers),
                "confidence": None if pd.isna(r.confidence) else float(r.confidence), "reason": r.reason or "", "stage": r.stage}
        if not any(p["doc"] == pair["doc"] and p["vor"] == pair["vor"] for p in pairs):
            pairs.append(pair)
    none_by_row = {(r.source_file, r.source_sheet, int(r.source_row)): {"confidence": float(r.confidence), "reason": r.reason or ""}
                   for r in none_rows.itertuples()}
    issues = issues.copy()
    ai_pairs, ai_none = [], []
    for _, r in issues.iterrows():
        counted = r["issue_type"] in ("volume_exceeded", "price_increase")
        ai_pairs.append(list(by_key.get(r["work_key"], [])) if counted else [])
        ai_none.append(none_by_row.get((r["source_file"], r["source_sheet"], int(r["source_row"]))) if r["issue_type"] == "missing_in_vor" else None)
    issues["ai_pairs"] = ai_pairs
    issues["ai_none"] = ai_none
    return issues


def collect_ai_examples(conn, project_id: str, cfg: dict) -> list:
    """Все решения ИИ из БД как примеры «правило кода -> ответ ИИ -> проверка кода» (отбор: pick_examples)."""
    m = cfg["rules"]["matching"]
    lo, hi = m["llm_lower_bound"], m["fuzzy_threshold"]
    min_conf = cfg["rules"]["llm"]["min_confidence"]
    labels = cfg["rules"]["issues"]["unit_labels"]
    markers = cfg["synonyms"].get("strip_markers", [])
    out = []
    for r in conn.execute(
            "SELECT c.candidate_id, c.score, c.status, c.judge_confidence, c.judge_reason, ia.work_name_raw AS doc_name, "
            "iv.work_name_raw AS vor_name, ga.unit_norm FROM staging_match_candidates c "
            "JOIN staging_match_groups ga ON ga.group_id = c.group_id JOIN staging_match_groups gv ON gv.group_id = c.candidate_group_id "
            "JOIN items ia ON ia.item_id = ga.first_item_id JOIN items iv ON iv.item_id = gv.first_item_id "
            "WHERE c.project_id = ? AND c.status IN ('accepted', 'rejected') ORDER BY c.candidate_id", (project_id,)):
        doc, vor, conf = _clean(r[5], markers), _clean(r[6], markers), r[3]
        unit = labels.get(r[7], r[7] or "")
        same = r[2] == "accepted"
        out.append({"kind": "pair_accepted" if same else "pair_rejected", "vor": vor, "doc": doc, "confidence": conf,
                    "code_rule": f"Схожесть названий {fmt_num(round(r[1]))} из 100: ниже порога автосклейки {hi}, зона проверки ИИ {lo}–{hi}",
                    "ai": (f"Одна и та же работа (уверенность {fmt_conf(conf)}): {r[4] or 'причина не указана'}" if same else
                           f"Разные работы (уверенность {fmt_conf(conf)}): {r[4] or 'причина не указана'}"),
                    "code_check": (f"Единица совпала ({unit}), уверенность {fmt_conf(conf)} не ниже порога {fmt_conf(min_conf)}, "
                                   "конфликтов по диаметру, классу и марке нет" if same else
                                   "Ответ «разные» принят: пара не склеена, строка осталась без пары")})
    for r in conn.execute(
            "SELECT d.decision_id, d.chosen_key, d.confidence, d.reason, d.outcome, ig.work_name_raw AS doc_name, "
            "(SELECT iv.work_name_raw FROM staging_match_groups gv JOIN items iv ON iv.item_id = gv.first_item_id "
            " WHERE gv.project_id = d.project_id AND gv.doc_type = 'vor' AND gv.final_work_key = d.chosen_key LIMIT 1) AS vor_name, "
            "(SELECT COUNT(*) FROM staging_match_candidates c WHERE c.group_id = d.group_id) AS n_cand, g.unit_norm "
            "FROM staging_llm_row_decisions d JOIN staging_match_groups g ON g.group_id = d.group_id "
            "JOIN items ig ON ig.item_id = g.first_item_id WHERE d.project_id = ? AND d.outcome IN ('accepted', 'none') "
            "ORDER BY d.decision_id", (project_id,)):
        if r[4] == "none" and r[2] < min_conf:
            continue
        doc, conf = _clean(r[5], markers), r[2]
        unit = labels.get(r[8], r[8] or "")
        rule = (f"Поиск по схожести названий ничего похожего не нашёл (ниже {lo} из 100)" if not r[7] else
                f"Нашлись только слабые кандидаты (схожесть ниже {hi} из 100): автоматически не решено")
        if r[4] == "accepted":
            vor = _clean(r[6], markers)
            out.append({"kind": "row_accepted", "vor": vor, "doc": doc, "confidence": conf, "code_rule": rule,
                        "ai": f"ИИ выбрал позицию ВОР из полного списка (уверенность {fmt_conf(conf)}): {r[3] or 'причина не указана'}",
                        "code_check": f"Позиция есть в списке ВОР, единица совпала ({unit}), числовых конфликтов нет, "
                                      "позицию не занимает другая строка акта"})
        else:
            out.append({"kind": "row_none", "vor": None, "doc": doc, "confidence": conf, "code_rule": rule,
                        "ai": f"В ВОР такой позиции нет (уверенность {fmt_conf(conf)}): {r[3] or 'причина не указана'}",
                        "code_check": f"Ответ принят: уверенность не ниже порога {fmt_conf(min_conf)}; расхождение подписано "
                                      f"«{AI_NOT_FOUND_NOTE}», решает специалист"})
    return out


def pick_examples(examples: list, n: int = 6) -> list:
    """Отбор для экрана: по порядку видов (2 склейки, 1 отказ, 2 выбора по списку, 1 «нет в ВОР»), без повторов названия в документе."""
    plan = [("pair_accepted", 2), ("pair_rejected", 1), ("row_accepted", 2), ("row_none", 1)]
    chosen, seen = [], set()
    for kind, want in plan:
        got = 0
        for e in examples:
            if e["kind"] == kind and e["doc"].lower() not in seen and got < want:
                chosen.append(e)
                seen.add(e["doc"].lower())
                got += 1
    for e in examples:                                  # не хватило какого-то вида: добираем любыми
        if len(chosen) >= n:
            break
        if e not in chosen and e["doc"].lower() not in seen:
            chosen.append(e)
            seen.add(e["doc"].lower())
    return chosen[:n]


def ai_pair_list(pair_rows: pd.DataFrame, markers) -> list:
    """Все пары названий, которые сопоставил ИИ (без повторов): [{doc_type, doc, vor, confidence, reason, stage}], сначала акты."""
    seen, out = set(), []
    for r in pair_rows.itertuples():
        pair = {"doc_type": r.doc_type, "doc": _clean(r.doc_name, markers), "vor": _clean(r.vor_name, markers),
                "confidence": None if pd.isna(r.confidence) else float(r.confidence), "reason": r.reason or "", "stage": r.stage}
        if (pair["doc"], pair["vor"]) not in seen:
            seen.add((pair["doc"], pair["vor"]))
            out.append(pair)
    return sorted(out, key=lambda p: (p["doc_type"] != "act", p["doc"].lower()))


def false_issue_ids(issues: pd.DataFrame) -> set:
    """issue_id расхождений, которых нет в эталоне синтетики (ложные): scripts/evaluate.py по data/ground_truth.csv."""
    scripts = str(ROOT / "scripts")
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    import evaluate as ev
    result = ev.evaluate_issues(issues.to_dict("records"), ev.read_csv(ROOT / "data" / "ground_truth.csv"), ev.read_csv(ROOT / "data" / "traps.csv"))
    return {int(i["issue_id"]) for i in result["false_issues"]}


def compact_rows(issues: pd.DataFrame, positions: pd.DataFrame, false_ids: set) -> pd.DataFrame:
    """Короткая таблица расхождений одного режима для страницы обоснования: тип, работа, влияние, совпало ли с эталоном."""
    titles = _titles(issues, positions)
    rows = []
    for idx, r in issues.iterrows():
        impact = EMPTY if pd.isna(r["impact_som"]) else f"{fmt_num(r['impact_som'])} сом"
        rows.append({"Тип": TYPE_RU[r["issue_type"]], "Работа": titles[idx], "Влияние": impact,
                     "По эталону": "ложное: в эталоне нет" if int(r["issue_id"]) in false_ids else "есть в эталоне"})
    return pd.DataFrame(rows, columns=["Тип", "Работа", "Влияние", "По эталону"])


def quality_metrics(issues: pd.DataFrame) -> dict:
    """Число расхождений, ложные расхождения, найдено из эталона и точность: считает scripts/evaluate.py по data/ground_truth.csv
    (эталон синтетического генератора). Точность: доля показанных расхождений, которые есть в эталоне."""
    scripts = str(ROOT / "scripts")
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    import evaluate as ev
    gt = ev.read_csv(ROOT / "data" / "ground_truth.csv")
    traps = ev.read_csv(ROOT / "data" / "traps.csv")
    result = ev.evaluate_issues(issues.to_dict("records"), gt, traps)
    return {"issues": result["issues"], "false": result["false_count"], "found": result["found"], "gt_total": result["gt_total"],
            "precision": result["precision"]}


def quality_compare(rules_issues: pd.DataFrame, llm_issues: pd.DataFrame) -> list:
    """Строки блока «Без ИИ и с ИИ»: [(показатель, без ИИ, с ИИ)] готовыми строками."""
    a, b = quality_metrics(rules_issues), quality_metrics(llm_issues)
    pct = lambda q: EMPTY if q["precision"] is None else fmt_pct(q["precision"] * 100, signed=False)  # noqa: E731
    return [("Возможных расхождений", fmt_num(a["issues"]), fmt_num(b["issues"])),
            ("Из них ложных", fmt_num(a["false"]), fmt_num(b["false"])),
            ("Найдено заложенных", f"{a['found']} из {a['gt_total']}", f"{b['found']} из {b['gt_total']}"),
            ("Точность", pct(a), pct(b))]


def detail_rows(without_ai: dict, with_ai: dict) -> list:
    """Подробное сравнение режимов: качество по эталону и числа из summary (строки без пары, светофор, позиции с отклонением от плана, влияние)."""
    rows = quality_compare(without_ai["issues"], with_ai["issues"])
    sa, sb = without_ai["summary"], with_ai["summary"]
    ma = headline_metrics(sa, without_ai["issues"], without_ai["positions"])
    mb = headline_metrics(sb, with_ai["issues"], with_ai["positions"])
    light = lambda s: " / ".join(str(s["statuses"][k]) for k in ("red", "yellow", "green"))  # noqa: E731
    impact_a = fmt_num(ma["impact"]) + (f" (и ещё {fmt_num(ma['impact_extra'])} низкой уверенности)" if ma["impact_extra_n"] else "")
    return rows + [("Строк без пары", fmt_num(sa["k"]), fmt_num(sb["k"])),
                   ("Строк без решения, нужна проверка", fmt_num(sa["a"]), fmt_num(sb["a"])),
                   ("Светофор: красные / жёлтые / зелёные", light(sa), light(sb)),
                   ("Позиции с отклонением от плана (красные и жёлтые)", fmt_num(ma["review_positions"]), fmt_num(mb["review_positions"])),
                   ("Возможное влияние на бюджет, сом", impact_a, fmt_num(mb["impact"]))]


# ---------- сводные числа ----------
def impact_split(issues: pd.DataFrame) -> dict:
    """Сумма влияния на бюджет по высокой и низкой уверенности (пустое влияние считается нулём) и число расхождений низкой уверенности."""
    if issues.empty:
        return {"high": 0.0, "low": 0.0, "n_low": 0}
    amount = issues["impact_som"].fillna(0)
    low = issues["confidence"] == "low"
    return {"high": float(amount[~low].sum()), "low": float(amount[low].sum()), "n_low": int(low.sum())}


def headline_metrics(summary: dict, issues: pd.DataFrame, positions: pd.DataFrame) -> dict:
    """Четыре карточки: позиций проверено, расхождений, влияние на бюджет, позиции с отклонением от плана (красные и жёлтые из position_status).

    Влияние: в режиме llm общая сумма из summary; в режиме без ИИ крупно только высокая уверенность, остальное отдельно (мелко)."""
    split = impact_split(issues)
    llm = summary["mode"] == "llm"
    review = int(positions["status"].isin(["red", "yellow"]).sum()) if not positions.empty else 0
    return {"positions": summary["n"], "issues": summary["z"],
            "impact": summary["impact_som"] if llm else split["high"],
            "impact_extra_n": 0 if llm else split["n_low"], "impact_extra": 0.0 if llm else split["low"],
            "review_positions": review}


def traffic_segments(summary: dict) -> list:
    """Три сегмента светофора из summary['statuses']: [{status, count, word, bg, fg, label}]."""
    out = []
    for status in ("red", "yellow", "green"):
        bg, fg = STATUS_COLORS[status]
        count = int(summary["statuses"][status])
        out.append({"status": status, "count": count, "word": STATUS_PLURAL[status], "bg": bg, "fg": fg,
                    "label": f"{STATUS_PLURAL[status]}: {count} {_positions_word(count)}"})
    return out


def _positions_word(n: int) -> str:
    if 11 <= n % 100 <= 14:
        return "позиций"
    return {1: "позиция", 2: "позиции", 3: "позиции", 4: "позиции"}.get(n % 10, "позиций")


def traffic_legend() -> str:
    """Пояснение цветов по docs/synthetic_spec.md §6; допуски берутся из config/rules.yaml, не хардкодятся."""
    rules = load_config()["rules"]
    over = rules["volume_exceeded"]["tolerance_pct"]
    under = rules["position_status"]["green_under_tolerance_pct"]
    return (f"Красный: есть возможное расхождение по позиции или факт по актам больше плана более чем на {fmt_num(over)} %. "
            f"Жёлтый: факт меньше плана более чем на {fmt_num(under)} % (работа идёт или пара ещё не найдена). "
            "Зелёный: факт в пределах допуска.")


def chart_frames(issues: pd.DataFrame) -> tuple:
    """-> (по типам: число расхождений, по типам: влияние на бюджет). Все четыре типа всегда присутствуют, подписи готовые строки."""
    counts, impacts = [], []
    for t in TYPE_ORDER:
        part = issues[issues["issue_type"] == t] if not issues.empty else issues
        n = int(len(part))
        amount = float(part["impact_som"].fillna(0).sum()) if n else 0.0
        counts.append({"label": TYPE_CHART_RU[t], "value": n, "text": str(n)})
        impacts.append({"label": TYPE_CHART_RU[t], "value": amount, "text": fmt_num(amount)})
    return pd.DataFrame(counts), pd.DataFrame(impacts)


# ---------- карточки расхождений ----------
def issue_phrase(r) -> str:
    """Одна понятная фраза «что не так» с числами из полей issues_view (expected, actual, delta_pct, unit, expected_text, actual_text)."""
    unit = r["unit"] or ""
    kind = r["issue_type"]
    if kind == "volume_exceeded":
        return (f"В актах {fmt_num(r['actual'])} {unit} при {fmt_num(r['expected'])} {unit} в ВОР ({fmt_pct(r['delta_pct'])})").replace("  ", " ")
    if kind == "price_increase":
        if pd.isna(r["expected"]):
            return f"Цена в акте {fmt_num(r['actual'])} сом за {unit}; цены в смете для сравнения нет"
        return (f"Цена в акте {fmt_num(r['actual'])} сом за {unit} при {fmt_num(r['expected'])} сом в смете "
                f"({fmt_pct(r['delta_pct'])})")
    if kind == "missing_in_vor":
        return f"Позиция есть в актах ({fmt_num(r['actual'])} {unit}), но в ВОР пары не найдено".replace("  ", " ")
    return (f"Акт датирован {fmt_date(r['actual_text'])}, срок договора {fmt_date(r['expected_text'])}: "
            f"позже срока на {fmt_num(r['actual'])} дн.")


def sort_issues(issues: pd.DataFrame) -> pd.DataFrame:
    """Сначала высокая важность, внутри по влиянию на бюджет по убыванию (пустое влияние в конце группы)."""
    if issues.empty:
        return issues
    rank = issues["severity"].map({s: i for i, s in enumerate(SEVERITY_ORDER)})
    keyed = issues.assign(_rank=rank, _impact=issues["impact_som"].fillna(-1))
    return keyed.sort_values(["_rank", "_impact", "issue_id"], ascending=[True, False, True]).drop(columns=["_rank", "_impact"])


def filter_issues(issues: pd.DataFrame, positions: pd.DataFrame, severity: str = "all", types=None, query: str = "") -> pd.DataFrame:
    """Фильтры экрана: важность (high/medium/low/all), типы (список issue_type или пусто = все), поиск по названию работы."""
    out = issues
    if severity != "all":
        out = out[out["severity"] == severity]
    if types:
        out = out[out["issue_type"].isin(types)]
    if query.strip():
        names = _titles(out, positions)
        mask = names.str.contains(query.strip(), case=False, regex=False)
        out = out[mask.values]
    return out


def _titles(issues: pd.DataFrame, positions: pd.DataFrame) -> pd.Series:
    names = dict(zip(positions["work_key"], positions["name"])) if not positions.empty else {}
    titles = []
    for _, r in issues.iterrows():
        if r["issue_type"] == "late_act":
            titles.append(f"Акт {r['source_file']} после срока договора")
        else:
            titles.append(names.get(r["work_key"]) or (r["work_key"] or EMPTY))
    return pd.Series(titles, index=issues.index, dtype="object")


def build_cards(issues: pd.DataFrame, positions: pd.DataFrame) -> list:
    """Карточки расхождений в порядке датафрейма (отсортируй sort_issues заранее)."""
    titles = _titles(issues, positions)
    cards = []
    for idx, r in issues.iterrows():
        if pd.isna(r["impact_som"]):
            impact = f"{EMPTY} ({r['impact_note']})" if r["impact_note"] else EMPTY
        else:
            impact = f"{fmt_num(r['impact_som'])} сом"
        if r["confidence"] == "low":
            note = LOW_CONFIDENCE_NOTE
        elif r["issue_type"] == "missing_in_vor":
            note = AI_NOT_FOUND_NOTE
        else:
            note = ""
        cards.append({"issue_id": int(r["issue_id"]), "title": titles[idx], "type": r["issue_type"], "type_label": TYPE_RU[r["issue_type"]],
                      "severity": r["severity"], "severity_label": SEVERITY_RU[r["severity"]], "phrase": issue_phrase(r),
                      "impact_text": impact, "impact_value": None if pd.isna(r["impact_som"]) else float(r["impact_som"]),
                      "sources": list(r["sources"]), "ai": bool(r["ai_matched"]), "note": note,
                      "ai_pairs": list(r["ai_pairs"]), "ai_none": r["ai_none"], "sides": r["sides"]})
    return cards


def cards_frame(cards: list) -> pd.DataFrame:
    """Те же карточки таблицей: влияние числом (для форматирования колонки), источники в одну колонку."""
    return pd.DataFrame([{"Важность": c["severity_label"], "Тип": c["type_label"], "Работа": c["title"], "Что не так": c["phrase"],
                          "Влияние, сом": c["impact_value"], "Источник": "\n".join(c["sources"]),
                          "ИИ помог сопоставить": "да" if c["ai"] else "", "Пометка": c["note"]} for c in cards],
                        columns=["Важность", "Тип", "Работа", "Что не так", "Влияние, сом", "Источник", "ИИ помог сопоставить", "Пометка"])


# ---------- светофор по позициям ----------
def positions_view(positions: pd.DataFrame, status: str = "all") -> pd.DataFrame:
    """Таблица позиций: красные сверху; колонка Статус словом. Служебная колонка _status нужна только для подсветки строк."""
    cols = ["Работа", "Ед.", "План", "Факт", "Выполнено, %", "Статус", "_status"]
    if positions.empty:
        return pd.DataFrame(columns=cols)
    out = positions if status == "all" else positions[positions["status"] == status]
    rank = out["status"].map({"red": 0, "yellow": 1, "green": 2})
    out = out.assign(_rank=rank).sort_values(["_rank", "name"])
    return pd.DataFrame({"Работа": out["name"].values, "Ед.": out["unit_label"].values, "План": out["plan_qty"].values,
                         "Факт": out["fact_qty"].values, "Выполнено, %": out["pct"].values,
                         "Статус": out["status"].map(STATUS_SINGULAR).values, "_status": out["status"].values}, columns=cols)
