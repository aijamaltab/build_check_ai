"""Отчёт по сверке: Excel-книга (.xlsx) с листами «Расхождения» и «Позиции» и кнопка скачивания."""
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

from ui.data import build_cards, cards_frame, positions_view, sort_issues


def generate_excel_report(issues: pd.DataFrame, positions: pd.DataFrame) -> bytes:
    """Генерация Excel-книги с двумя листами: «Расхождения» и «Позиции»."""
    cards = build_cards(sort_issues(issues), positions)
    df_issues = cards_frame(cards)
    df_pos = positions_view(positions)

    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="E8ECEF", end_color="E8ECEF", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True)
    regular_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)

    thin_border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE")
    )

    # ---------- Лист 1: Расхождения ----------
    ws_issues = wb.active
    ws_issues.title = "Расхождения"

    cols_issues = ["№", "Важность", "Тип", "Работа", "Что не так", "Влияние, сом", "Источник"]
    ws_issues.append(cols_issues)
    for col_idx in range(1, len(cols_issues) + 1):
        cell = ws_issues.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, (_, r) in enumerate(df_issues.iterrows(), 2):
        ws_issues.append([
            row_idx - 1,
            r["Важность"],
            r["Тип"],
            r["Работа"],
            r["Что не так"],
            r["Влияние, сом"],
            r["Источник"]
        ])
        for col_idx in range(1, len(cols_issues) + 1):
            cell = ws_issues.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in (2, 3):
                cell.alignment = Alignment(horizontal="center")
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0.00"

    # ---------- Лист 2: Позиции ----------
    ws_pos = wb.create_sheet(title="Позиции")
    cols_pos = ["№", "Работа", "Ед.", "План", "Факт", "Выполнено, %", "Статус"]
    ws_pos.append(cols_pos)
    for col_idx in range(1, len(cols_pos) + 1):
        cell = ws_pos.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    status_fills = {
        "red": PatternFill(start_color="FDECEA", fill_type="solid"),
        "yellow": PatternFill(start_color="FFF6DB", fill_type="solid"),
        "green": PatternFill(start_color="E8F5E9", fill_type="solid")
    }

    for row_idx, (_, r) in enumerate(df_pos.iterrows(), 2):
        ws_pos.append([
            row_idx - 1,
            r["Работа"],
            r["Ед."],
            r["План"],
            r["Факт"],
            r["Выполнено, %"],
            r["Статус"]
        ])
        st_val = r.get("_status", "")
        fill_to_use = status_fills.get(st_val)

        for col_idx in range(1, len(cols_pos) + 1):
            cell = ws_pos.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if fill_to_use and col_idx == 7:
                cell.fill = fill_to_use
                cell.font = bold_font
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in (4, 5, 6):
                cell.alignment = Alignment(horizontal="right")
                if col_idx in (4, 5):
                    cell.number_format = "#,##0.00"
                elif col_idx == 6:
                    cell.number_format = "#,##0.0"
            elif col_idx in (3, 7):
                cell.alignment = Alignment(horizontal="center")

    # Автоподбор ширины колонок
    for ws in (ws_issues, ws_pos):
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 65)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def download_button(issues: pd.DataFrame, positions: pd.DataFrame, key: str = "report") -> None:
    st.download_button(label="Скачать отчёт в Excel (.xlsx)", data=generate_excel_report(issues, positions), file_name="buildcheck_report.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary", icon=":material/download:", key=key)

