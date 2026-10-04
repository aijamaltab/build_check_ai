"""Маленькие проекты из настоящих xlsx по шаблонам vor_a, estimate_a, act_a, act_c и contract_a для тестов проверок.

Файлы собираются openpyxl с теми же заголовками, что в config/templates.yaml, и проходят через обычный ingestion.
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook


def _save(wb, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def make_project(folder, vor, estimate, acts, deadline=dt.date(2025, 9, 30), act_dates=None, usd_acts=()):
    """vor: [(название, ед., количество)]; estimate: [(название, ед., количество, цена за единицу)];
    acts: список актов, каждый [(название, ед., количество, цена или None)]; номера актов с 1.
    act_dates: даты актов (по умолчанию до срока договора); usd_acts: номера актов в шаблоне act_c с валютой USD в заголовках."""
    folder = Path(folder)
    wb = Workbook()
    ws = wb.active
    ws.title = "ВОР"
    ws.append(["ВЕДОМОСТЬ ОБЪЁМОВ РАБОТ"])
    ws.append([])
    ws.append(["№ п.п", "Наименование работ и затрат", "Ед. изм.", "Кол-во", "Формула расчёта объёма"])
    for i, (name, unit, qty) in enumerate(vor, 1):
        ws.append([i, name, unit, qty, None])
    _save(wb, folder / "vor.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    ws.append(["ЛОКАЛЬНАЯ СМЕТА № 1"])
    ws.append(["№ поз.", "Шифр норматива", "Наименование", "Ед. изм.", "Кол-во", "Стоимость единицы, сом", "Общая стоимость, сом"])
    total = 0
    for i, (name, unit, qty, price) in enumerate(estimate, 1):
        ws.append([i, None, name, unit, qty, price, qty * price])
        total += qty * price
    ws.append([None, None, "Итого по смете, сом", None, None, None, total])
    _save(wb, folder / "estimate.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Договор"
    ws.append(["ДОГОВОР ПОДРЯДА"])
    ws.append(["Срок выполнения работ до", deadline])
    ws.append(["Цена договора, сом", 1000000])
    _save(wb, folder / "contract.xlsx")

    for n, rows in enumerate(acts, 1):
        date = (act_dates or {}).get(n, dt.date(2025, 5, 31))
        wb = Workbook()
        ws = wb.active
        ws.title = "Акт"
        total = 0
        if n in usd_acts:                                   # шаблон act_c: дата ячейкой над шапкой, валюта в заголовках
            ws.append(["Акт приёмки выполненных работ"])
            ws.append([None, None, date])
            ws.append([])
            ws.append(["№", "Наименование работ", "Единица измерения", "Количество", "Цена за единицу, USD", "Сумма, USD", "Примечание"])
            for i, (name, unit, qty, price) in enumerate(rows, 1):
                ws.append([i, name, unit, qty, price, None if price is None else qty * price, None])
                total += 0 if price is None else qty * price
        else:
            ws.append([f"АКТ № {n} от {date.strftime('%d.%m.%Y')} о приёмке выполненных работ"])
            ws.append([])
            ws.append(["№", "Наименование выполненных работ", "Ед.", "Выполнено", "Цена", "Сумма"])
            for i, (name, unit, qty, price) in enumerate(rows, 1):
                ws.append([i, name, unit, qty, price, None if price is None else qty * price])
                total += 0 if price is None else qty * price
            ws.append([None, "Итого по акту, сом", None, None, None, total])
        _save(wb, folder / f"act_{n}.xlsx")
    return folder
