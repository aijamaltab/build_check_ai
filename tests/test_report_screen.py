"""Тесты страницы «Скачать отчёт» (ui/screens/report.py) и генератора Excel-файла."""
import io
import tempfile
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from tests.cache_guard import REAL_CACHE
from ui.loader import build_demo
from ui.screens.report import generate_excel_report

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture()
def llm_env(monkeypatch, tmp_path):
    monkeypatch.setenv("LLM_CACHE_DIR", str(REAL_CACHE))
    monkeypatch.setenv("LLM_CACHE_ONLY", "1")
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))


def report_app():
    import os
    os.environ.setdefault("LLM_CACHE_ONLY", "1")
    from ui.screens import report
    report.render()


def test_generate_excel_report_structure(llm_env):
    data = build_demo("llm")
    excel_bytes = generate_excel_report(data["issues"], data["positions"])
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert wb.sheetnames == ["Расхождения", "Позиции"]

    # Лист Расхождения
    ws_issues = wb["Расхождения"]
    assert ws_issues.max_row == 14  # 1 заголовок + 13 расхождений
    headers_issues = [cell.value for cell in ws_issues[1]]
    assert headers_issues == ["№", "Важность", "Тип", "Работа", "Что не так", "Влияние, сом", "Источник"]

    # Лист Позиции
    ws_pos = wb["Позиции"]
    assert ws_pos.max_row == 49  # 1 заголовок + 48 позиций
    headers_pos = [cell.value for cell in ws_pos[1]]
    assert headers_pos == ["№", "Работа", "Ед.", "План", "Факт", "Выполнено, %", "Статус"]


def test_report_page_render_and_download_button(llm_env):
    at = AppTest.from_function(report_app, default_timeout=60).run()
    assert not at.exception

    html = "\n".join(m.value for m in at.markdown)
    assert "Скачать отчёт о сверке" in html
    assert "Экспорт данных" in html

    # Проверяем метрики
    assert "13" in html
    assert "1 378 030" in html

    # Кнопка скачивания
    assert len(at.download_button) >= 1
    btn = at.download_button[0]
    assert "Скачать отчёт в Excel" in btn.label

    # Вкладки предпросмотра
    assert len(at.tabs) == 2
    assert "«Расхождения» (13)" in at.tabs[0].label
    assert "«Позиции» (48)" in at.tabs[1].label
