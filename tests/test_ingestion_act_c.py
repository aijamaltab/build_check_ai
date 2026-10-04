"""Шаблон act_c на маленькой книге в памяти с ВЫМЫШЛЕННЫМИ нейтральными данными.

Книга повторяет структуру из docs/real_templates_findings.md (раздел 4), но названия, суммы и объект выдуманы:
шапка не в первой строке, дата отдельной ячейкой над шапкой, строка объекта без номера, валюта USD только
в заголовках цены и суммы, цена формулой «сумма / количество», безымянные служебные колонки справа,
одно название в нескольких строках.
"""
import datetime as dt
import io
import re
import zipfile

from openpyxl import Workbook

from src.config import load_config
from src.ingestion.items import build_items
from src.ingestion.parse import parse_workbook

CFG = load_config()


def inject_cached_values(data: bytes, values: dict) -> bytes:
    """openpyxl пишет формулы без кэшированного значения. Как в файле из Excel, подставляем <v> вручную."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            blob = zin.read(name)
            if name == "xl/worksheets/sheet1.xml":
                text = blob.decode("utf-8")
                for ref, val in values.items():
                    text, n = re.subn(rf'(<c r="{ref}"[^>]*><f>[^<]*</f>)(?:<v\s*/>|<v></v>)', rf"\g<1><v>{val}</v>", text)
                    assert n == 1, f"не нашли формулу в {ref}"
                blob = text.encode("utf-8")
            zout.writestr(name, blob)
    return buf.getvalue()


def build_book() -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Акт 7"
    ws["A1"] = "Акт приёмки выполненных работ"
    ws["C2"] = dt.datetime(2025, 6, 15)                       # дата отдельной ячейкой, без подписи
    headers = ["№", "Наименование работ", "Единица измерения", "Количество", "Цена за единицу, USD", "Сумма, USD",
               "Примечание", None, None]
    for c, h in enumerate(headers, 1):
        ws.cell(5, c, h)                                       # шапка в строке 5, не в первой
    ws["A6"] = "Условный объект для теста (раздел 1)"          # строка объекта без номера
    ws.merge_cells("A6:G6")
    rows = [  # номер, название, единица, количество, сумма
        (1, "Условная работа Альфа", "м2", 10, 1000),
        (2, "Условная работа Альфа", "м2", 4, 400),            # то же название, другое количество: не дубль
        (3, "Условная работа Бета", "шт", 3, 900),
        (4, "Условная работа Гамма", "м", 5, 250),             # цена-формула без кэшированного значения
    ]
    for i, (no, name, unit, qty, amount) in enumerate(rows):
        r = 7 + i
        ws.cell(r, 1, no), ws.cell(r, 2, name), ws.cell(r, 3, unit), ws.cell(r, 4, qty)
        ws.cell(r, 5, f"=F{r}/D{r}")                           # цена = сумма / количество
        ws.cell(r, 6, amount)
        ws.cell(r, 7, "100%")
        ws.cell(r, 8, f"=F{r}*0.5")                            # безымянные служебные колонки справа
        ws.cell(r, 9, 12345)
    ws["B11"], ws["F11"] = "Итого", 2550
    buf = io.BytesIO()
    wb.save(buf)
    # кэшированные значения цены у первых трёх строк; у четвёртой нет (как у файла, не пересчитанного в Excel)
    return io.BytesIO(inject_cached_values(buf.getvalue(), {"E7": 100, "E8": 100, "E9": 300}))


def parse():
    return parse_workbook(build_book(), "act_c_test.xlsx", CFG)


def test_act_c_header_found_by_words_and_date_above_header():
    doc = parse()
    assert (doc.template, doc.doc_type, doc.errors) == ("act_c", "act", [])
    assert doc.doc_date == dt.date(2025, 6, 15) and doc.date_row == 2


def test_act_c_currency_comes_from_header_text():
    assert parse().currency == "USD"


def test_act_c_object_row_unnamed_columns_and_totals_are_not_items():
    doc = parse()
    names = [r.name for r in doc.rows]
    assert "Условный объект для теста (раздел 1)" not in names
    assert all("итого" not in n.lower() for n in names)
    assert doc.total_amount == 2550 and doc.total_row == 11


def test_act_c_repeated_names_are_separate_rows():
    rows = [r for r in parse().rows if r.name == "Условная работа Альфа"]
    assert [(r.quantity, r.amount) for r in rows] == [(10, 1000), (4, 400)]


def test_act_c_price_formula_uses_cached_value_and_is_flagged():
    doc = parse()
    alpha = doc.rows[0]
    assert alpha.unit_price == 100 and alpha.price_is_formula
    items = build_items(doc, CFG)
    assert items[0].currency == "USD" and items[0].doc_date == "2025-06-15" and items[0].price_is_formula


def test_act_c_formula_without_cached_value_is_not_counted_silently():
    doc = parse()
    assert doc.formula_missing == [10]
    assert [r for r, _ in doc.unrecognized] == [10]
    assert len(doc.rows) == 3
