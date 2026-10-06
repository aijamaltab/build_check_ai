"""Дополнительные наборы синтетики (data/synthetic_sets): воспроизводимость, эталон, ingestion, отличие от базового набора."""
import csv
import importlib.util
import sys
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.pipeline import run_pipeline

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import generate_synthetic as gen  # noqa: E402
from synthetic_profiles import PROFILES, set_dir  # noqa: E402

SETS = ["set_2", "set_3", "set_4"]
XLSX = ["vor_1", "vor_2", "estimate", "contract", "act_1", "act_2", "act_3", "act_4", "act_5"]


def norm(path) -> bytes:
    """Байты файла; у CSV переводы строк приводятся к LF (git на Windows может подставить CRLF при checkout)."""
    data = Path(path).read_bytes()
    return data.replace(bytes([13, 10]), bytes([10])) if str(path).endswith(".csv") else data


def cells(path):
    return [[c.value for c in row] for row in load_workbook(path).active.iter_rows()]


def variants(profile):
    with gen.profile_context(profile):
        return set(gen.all_name_variants()) - {gen.normalize_name(gen.WASTE_NAME)}      # «Строительный мусор» общий для ловушки 5


def test_base_set_is_generated_exactly_as_before(tmp_path):
    gen.generate(out_dir=tmp_path, meta_dir=tmp_path)
    for name in XLSX:
        assert cells(tmp_path / f"{name}.xlsx") == cells(ROOT / "data" / "synthetic" / f"{name}.xlsx")     # содержимое старого набора не изменилось
    for name in ("ground_truth.csv", "traps.csv", "expected_status.csv"):
        assert norm(tmp_path / name) == norm(ROOT / "data" / name)


def test_profile_does_not_leak_into_module_state(tmp_path):
    before = (gen.ITEMS[0].name_a, dict(gen.OBJECT), gen.TRAPS[0][2])
    gen.generate(PROFILES["set_2"]["seed"], tmp_path, tmp_path, PROFILES["set_2"])
    assert (gen.ITEMS[0].name_a, dict(gen.OBJECT), gen.TRAPS[0][2]) == before


@pytest.mark.parametrize("name", SETS)
def test_committed_set_matches_the_generator_byte_for_byte(name, tmp_path):
    profile = PROFILES[name]
    gen.generate(profile["seed"], tmp_path, tmp_path, profile)
    for path in sorted(tmp_path.iterdir()):
        assert norm(path) == norm(set_dir(name) / path.name), path.name


@pytest.mark.parametrize("name", SETS)
def test_set_has_nine_files_with_ground_truth_and_traps(name):
    folder = set_dir(name)
    assert sorted(p.stem for p in folder.glob("*.xlsx")) == sorted(XLSX)
    with open(folder / "ground_truth.csv", encoding="utf-8", newline="") as f:
        gt = list(csv.DictReader(f))
    with open(folder / "traps.csv", encoding="utf-8", newline="") as f:
        traps = list(csv.DictReader(f))
    assert len(gt) == 12 and len(traps) == 6
    assert sorted(g["issue_type"] for g in gt).count("volume_exceeded") == 5
    assert {g["source_file"] for g in gt} <= {p.name for p in folder.glob("*.xlsx")}


@pytest.mark.parametrize("name", SETS)
def test_set_passes_ingestion_without_errors(name, tmp_path):
    summary = run_pipeline(set_dir(name), tmp_path / "x.db", "rules_only")
    assert summary["files"] == 9 and summary["documents_with_errors"] == [] and summary["m"] == 0


def test_sets_use_other_names_objects_and_seeds_than_the_base_set():
    names = {"base": variants(None), **{n: variants(PROFILES[n]) for n in SETS}}
    keys = list(names)
    for i, x in enumerate(keys):
        for y in keys[i + 1:]:
            assert not names[x] & names[y], (x, y)                                                   # ни одного общего названия работы
    assert len({PROFILES[n]["seed"] for n in SETS} | {gen.DEFAULT_SEED}) == len(SETS) + 1
    assert len({str(PROFILES[n]["object"]) for n in SETS}) == len(SETS) and "СШ № 99" not in str([PROFILES[n]["object"] for n in SETS])


@pytest.mark.parametrize("name", SETS)
def test_trap_3_description_names_the_real_act_wording(name):
    import re
    with open(set_dir(name) / "traps.csv", encoding="utf-8", newline="") as f:
        text = next(r["description"] for r in csv.DictReader(f) if r["trap_id"] == "3")
    act_name = re.search(r"в акте «([^»]+)»", text).group(1)
    acts = " ".join(str(v) for n in ("act_1", "act_2") for row in cells(set_dir(name) / f"{n}.xlsx") for v in row if v)
    assert act_name in acts                                                                           # описание ловушки совпадает с данными


@pytest.mark.parametrize("name", SETS)
def test_set_text_is_synthetic_and_has_no_real_data_markers(name):
    texts = " ".join(str(v) for path in set_dir(name).glob("*.xlsx") for row in cells(path) for v in row if v is not None)
    assert "синтетическ" in texts and "вымышленн" in texts
    for marker in ("ИНН", "ОГРН", "@", "http", "+996"):
        assert marker not in texts


def test_scaling_keeps_shares_of_the_embedded_discrepancies():
    with open(set_dir("set_2") / "ground_truth.csv", encoding="utf-8", newline="") as f:
        gt = {g["gt_id"]: g for g in csv.DictReader(f)}
    assert round(float(gt["2"]["actual"]) / float(gt["2"]["expected"]), 6) == round(11.4 / 9.5, 6)     # арматура: +20 % как в базовом наборе
    assert round(float(gt["7"]["actual"]) / (6200 * 1.2), 6) == round(6900 / 6200, 6)  # кладка: цена в акте +11,3 % к смете
